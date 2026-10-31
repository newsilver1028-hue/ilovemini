from unittest.mock import patch
from io import BytesIO
from urllib.parse import parse_qs, urlparse
import re

from django.contrib.auth import get_user_model
from django.test import Client, RequestFactory, TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from .models import (
    LedgerEntry,
    NaverIdentity,
    Partner,
    PartnerBooking,
    PartnerReview,
    PartnerStaff,
    PartnerAccessRequest,
    Vehicle,
    WebPushSubscription,
)
from .partner_admin import (
    PartnerBookingAdmin,
    PartnerLedgerEntryAdmin,
    partner_admin_site,
)


class PartnerPortalTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.user = User.objects.create_user(username="partner-user", password=None, is_staff=False)
        NaverIdentity.objects.create(user=self.user, subject="naver-subject-1", nickname="업체 담당자")
        self.partner = Partner.objects.create(name="테스트 업체", is_active=True)
        PartnerStaff.objects.create(
            user=self.user,
            partner=self.partner,
            can_manage_bookings=True,
            can_verify_records=True,
            is_active=True,
        )
        self.other_partner = Partner.objects.create(name="다른 업체", is_active=True)
        self.other_user = User.objects.create_user(username="other", password=None)
        self.other_vehicle = Vehicle.objects.create(model_name="MINI Cooper", generation="F56", model_year=2022)
        self.own_vehicle = Vehicle.objects.create(model_name="MINI Cooper S", generation="F56", model_year=2023, plate_number="12가3456")
        from django.utils import timezone
        self.booking = PartnerBooking.objects.create(
            customer=self.other_user,
            partner=self.partner,
            vehicle=self.own_vehicle,
            scheduled_at=timezone.now(),
            service_type="정비",
            contact_phone="010-1234-5678",
        )
        self.other_booking = PartnerBooking.objects.create(
            customer=self.other_user,
            partner=self.other_partner,
            vehicle=self.other_vehicle,
            scheduled_at=timezone.now(),
            service_type="타이어",
        )
        self.request_factory = RequestFactory()

    def request_for_partner(self):
        request = self.request_factory.get("/partner-admin/")
        request.user = self.user
        return request

    def test_partner_naver_oauth_uses_existing_identity_and_lands_in_scoped_admin(self):
        client = Client()
        with override_settings(
            NAVER_CLIENT_ID="client-id",
            NAVER_CLIENT_SECRET="client-secret",
            NAVER_REDIRECT_URI="https://example.test/api/auth/naver/callback/",
        ):
            response = client.get(reverse("partner_naver_start"), {"partner_id": self.partner.pk})
            self.assertEqual(response.status_code, 302)
            self.assertIn("nid.naver.com/oauth2.0/authorize", response["Location"])
            state = parse_qs(urlparse(response["Location"]).query)["state"][0]
            callback_client = Client()
            with patch("core.auth_views.exchange_code", return_value="access-token"), patch(
                "core.auth_views.get_profile", return_value={"id": "naver-subject-1", "nickname": "업체 담당자"}
            ):
                response = callback_client.get(reverse("naver_callback") + f"?state={state}&code=authorization-code")

        self.assertEqual(response.status_code, 200)
        match = re.search(rb"name='ticket' value='([^']+)'", response.content)
        self.assertIsNotNone(match)
        ticket = match.group(1).decode()
        response = client.post(
            reverse("partner_naver_complete"), {"ticket": ticket},
            HTTP_HOST="admin.ilovemini.co.kr",
        )
        self.assertRedirects(response, reverse("partner_admin:index"), fetch_redirect_response=False)
        self.assertEqual(str(client.session["_auth_user_id"]), str(self.user.pk))
        self.assertEqual(client.get(reverse("partner_admin:index"), HTTP_HOST="admin.ilovemini.co.kr").status_code, 200)
        replay = Client().post(
            reverse("partner_naver_complete"), {"ticket": ticket},
            HTTP_HOST="admin.ilovemini.co.kr",
        )
        self.assertRedirects(
            replay,
            "https://admin.ilovemini.co.kr/partner-admin/login/?naver_error=expired",
            fetch_redirect_response=False,
        )

    def test_oauth_does_not_create_unregistered_partner_users(self):
        client = Client()
        with override_settings(
            NAVER_CLIENT_ID="client-id",
            NAVER_CLIENT_SECRET="client-secret",
            NAVER_REDIRECT_URI="https://example.test/api/auth/naver/callback/",
        ):
            start_response = client.get(reverse("partner_naver_start"), {"partner_id": self.partner.pk})
            state = parse_qs(urlparse(start_response["Location"]).query)["state"][0]
            with patch("core.auth_views.exchange_code", return_value="access-token"), patch(
                "core.auth_views.get_profile", return_value={"id": "unknown-subject"}
            ):
                response = client.get(reverse("naver_callback") + f"?state={state}&code=authorization-code")
        self.assertRedirects(
            response,
            "https://admin.ilovemini.co.kr/partner-admin/login/?naver_error=not_partner",
            fetch_redirect_response=False,
        )
        self.assertFalse(PartnerAccessRequest.objects.filter(
            partner=self.partner, user__naver_identity__subject="unknown-subject",
        ).exists())

    def test_partner_staff_sees_only_own_partner_and_cannot_delete(self):
        request = self.request_for_partner()
        self.assertTrue(partner_admin_site.has_permission(request))
        booking_admin = PartnerBookingAdmin(PartnerBooking, partner_admin_site)
        ledger_admin = PartnerLedgerEntryAdmin(LedgerEntry, partner_admin_site)
        self.assertEqual(list(booking_admin.get_queryset(request)), [self.booking])
        self.assertFalse(booking_admin.has_delete_permission(request))
        self.assertFalse(ledger_admin.has_delete_permission(request))
        self.assertFalse(booking_admin.has_add_permission(request))
        self.assertFalse(ledger_admin.has_change_permission(request))
        self.assertEqual(
            {model._meta.model_name for model in partner_admin_site._registry},
            {"partnerbooking", "ledgerentry", "maintenancerecommendation", "serviceestimate"},
        )

    def test_direct_delete_url_is_denied_for_partner_user(self):
        client = Client()
        client.force_login(self.user)
        response = client.get(f"/partner-admin/core/partnerbooking/{self.booking.pk}/delete/")
        self.assertEqual(response.status_code, 403)
        self.assertTrue(PartnerBooking.objects.filter(pk=self.booking.pk).exists())

    def test_repair_history_is_read_only_and_partner_scoped(self):
        own_entry = LedgerEntry.objects.create(
            vehicle=self.own_vehicle,
            kind=LedgerEntry.Kind.SERVICE,
            source=LedgerEntry.Source.PARTNER,
            partner=self.partner,
            amount_krw=120000,
            description="정비 내역",
        )
        LedgerEntry.objects.create(
            vehicle=self.other_vehicle,
            kind=LedgerEntry.Kind.SERVICE,
            source=LedgerEntry.Source.PARTNER,
            partner=self.other_partner,
        )
        request = self.request_for_partner()
        ledger_admin = PartnerLedgerEntryAdmin(LedgerEntry, partner_admin_site)
        self.assertEqual(list(ledger_admin.get_queryset(request)), [own_entry])
        self.assertFalse(ledger_admin.has_view_permission(request, LedgerEntry.objects.filter(partner=self.other_partner).first()))

        client = Client()
        client.force_login(self.user)
        response = client.get(reverse("partner_admin:core_ledgerentry_changelist"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "차량번호")
        self.assertContains(response, "12가3456")
        self.assertContains(response, "정비기록 미니 캘린더")
        self.assertContains(response, "인증 정비기록 조회")
        self.assertNotContains(response, "보려면 기록을 선택")

    def test_inactive_membership_is_denied(self):
        PartnerStaff.objects.filter(user=self.user, partner=self.partner).update(is_active=False)
        self.assertFalse(partner_admin_site.has_permission(self.request_for_partner()))

    def test_partner_and_regular_admin_login_pages_are_separate(self):
        client = Client()
        partner_page = client.get(reverse("partner_admin:login"))
        regular_page = client.get(reverse("admin:login"))
        self.assertEqual(partner_page.status_code, 200)
        self.assertEqual(regular_page.status_code, 200)
        self.assertContains(partner_page, "네이버 아이디로 협력업체 로그인")
        self.assertContains(regular_page, "협력업체 네이버 로그인")
        self.assertContains(regular_page, 'id="login-form"')

    def test_booking_calendar_renders_month_view_and_keeps_partner_scope(self):
        client = Client()
        client.force_login(self.user)
        year, month = self.booking.scheduled_at.strftime("%Y-%m").split("-")
        response = client.get(
            reverse("partner_admin:core_partnerbooking_changelist"),
            {"scheduled_at__year": year, "scheduled_at__month": month},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "월간 예약 캘린더")
        self.assertContains(response, "정비")
        self.assertContains(response, "12가3456")
        self.assertContains(response, "010-1234-5678")
        self.assertNotContains(response, "naver-subject-1")
        self.assertContains(response, "예약대장 엑셀 다운로드")
        self.assertNotContains(response, "다른 업체")

    def test_calendar_status_filter_still_detects_conflicts_with_hidden_statuses(self):
        from datetime import timedelta

        confirmed = PartnerBooking.objects.create(
            customer=self.other_user,
            partner=self.partner,
            vehicle=self.own_vehicle,
            scheduled_at=self.booking.scheduled_at + timedelta(minutes=30),
            duration_minutes=60,
            service_type="확정 정비",
            contact_phone="010-1234-5678",
            status=PartnerBooking.Status.CONFIRMED,
        )
        client = Client()
        client.force_login(self.user)
        year, month = self.booking.scheduled_at.strftime("%Y-%m").split("-")
        response = client.get(
            reverse("partner_admin:core_partnerbooking_changelist"),
            {"scheduled_at__year": year, "scheduled_at__month": month, "status": PartnerBooking.Status.REQUESTED},
        )
        self.assertEqual(response.status_code, 200)
        visible_events = [
            event
            for week in response.context["partner_calendar_weeks"]
            for cell in week
            for event in cell["events"]
        ]
        self.assertEqual(len(visible_events), 1)
        self.assertTrue(visible_events[0]["has_conflict"])
        self.assertNotIn(confirmed.service_type, str(visible_events))

    def test_partner_booking_export_is_xlsx_and_only_contains_assigned_partner_rows(self):
        client = Client()
        client.force_login(self.user)
        response = client.get(reverse("partner_admin:partner-booking-export"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertTrue(response.content.startswith(b"PK"))
        rows = list(load_workbook(BytesIO(response.content), read_only=True).active.values)
        self.assertEqual(len(rows), 2)
        self.assertIn("테스트 업체", rows[1])
        self.assertIn("12가3456", rows[1])
        self.assertIn("010-1234-5678", rows[1])
        self.assertNotIn("naver-subject-1", rows[1])
        self.assertNotIn("partner-user", rows[1])
        self.assertNotIn("다른 업체", rows[1])

    def test_partner_booking_change_form_uses_korean_labels_without_customer_login_id(self):
        client = Client()
        client.force_login(self.user)
        response = client.get(reverse("partner_admin:core_partnerbooking_change", args=[self.booking.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "차량번호")
        self.assertContains(response, "12가3456")
        self.assertContains(response, "연락처")
        self.assertContains(response, "010-1234-5678")
        self.assertContains(response, "예약 상태")
        self.assertContains(response, 'id="id_status"')
        self.assertNotContains(response, ">History</a>")
        self.assertNotContains(response, "Customer:")
        self.assertNotContains(response, "naver-subject-1")

    def test_authorized_partner_booking_staff_can_register_web_push(self):
        client = Client()
        client.force_login(self.user)
        response = client.post(
            reverse("partner_web_push_subscribe"),
            data={
                "endpoint": "https://push.example.test/subscription/abc",
                "keys": {"p256dh": "public-key", "auth": "auth-secret"},
            },
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        subscription = WebPushSubscription.objects.get()
        self.assertEqual(subscription.user, self.user)
        self.assertEqual(subscription.partner, self.partner)

    def test_new_booking_sends_web_push_even_without_mobile_fcm(self):
        WebPushSubscription.objects.create(
            user=self.user,
            partner=self.partner,
            endpoint="https://push.example.test/subscription/abc",
            p256dh="public-key",
            auth="auth-secret",
        )
        from django.test import override_settings
        with override_settings(WEB_PUSH_VAPID_PRIVATE_KEY="test-private-key"):
            with patch("pywebpush.webpush") as webpush:
                from .push import send_booking_notification
                sent = send_booking_notification(self.booking)
        self.assertEqual(sent, 1)
        self.assertEqual(webpush.call_count, 1)
        self.assertEqual(webpush.call_args.kwargs["subscription_info"]["endpoint"], "https://push.example.test/subscription/abc")

    def test_partner_without_booking_permission_cannot_register_web_push(self):
        self.user.partner_memberships.update(can_manage_bookings=False)
        client = Client()
        client.force_login(self.user)
        response = client.post(
            reverse("partner_web_push_subscribe"),
            data={
                "endpoint": "https://push.example.test/subscription/abc",
                "keys": {"p256dh": "public-key", "auth": "auth-secret"},
            },
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 403)

    def test_browser_admin_changelist_renders_date_mini_calendar(self):
        operator = get_user_model().objects.create_superuser("operator", "operator@example.test", "test-password")
        PartnerReview.objects.create(partner=self.partner, user=self.other_user, rating=5, comment="좋아요")
        client = Client()
        client.force_login(operator)
        response = client.get(reverse("admin:core_partnerreview_changelist"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "날짜별 조회")
        self.assertContains(response, "1건")

    def test_partner_maintenance_export_is_scoped_and_contains_itemized_amounts(self):
        client = Client()
        client.force_login(self.user)
        LedgerEntry.objects.create(
            vehicle=self.own_vehicle,
            kind=LedgerEntry.Kind.SERVICE,
            source=LedgerEntry.Source.PARTNER,
            partner=self.partner,
            amount_krw=120000,
            line_items=[{"name": "엔진오일", "amount_krw": 120000}],
            description="정기 점검",
        )
        LedgerEntry.objects.create(
            vehicle=self.other_vehicle,
            kind=LedgerEntry.Kind.SERVICE,
            source=LedgerEntry.Source.PARTNER,
            partner=self.other_partner,
            amount_krw=50000,
        )
        response = client.get(reverse("partner_admin:partner-maintenance-export"))
        self.assertEqual(response.status_code, 200)
        rows = list(load_workbook(BytesIO(response.content), read_only=True).active.values)
        self.assertEqual(len(rows), 2)
        self.assertIn("테스트 업체", rows[1])
        self.assertIn("엔진오일: 120,000원", rows[1])
        self.assertNotIn("다른 업체", rows[1])

    def test_excel_exports_require_the_matching_partner_permission(self):
        PartnerStaff.objects.filter(user=self.user).update(can_manage_bookings=False)
        client = Client()
        client.force_login(self.user)
        self.assertEqual(client.get(reverse("partner_admin:partner-booking-export")).status_code, 404)
        self.assertEqual(client.get(reverse("partner_admin:partner-maintenance-export")).status_code, 200)
