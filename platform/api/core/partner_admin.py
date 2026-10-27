from django.contrib import admin
from django.contrib import messages
from django.contrib.admin import AdminSite
from django import forms
from django.db.models import F, Q, Sum
from django.db import transaction
from django.core.exceptions import PermissionDenied, ValidationError
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, render
from django.urls import path, reverse
from django.utils.html import format_html_join
from django.utils import timezone
from django.utils.decorators import method_decorator
from django.views.decorators.cache import never_cache
from django.conf import settings
from datetime import datetime, time, timedelta

from .models import (
    LedgerEntry,
    LedgerEvidence,
    Partner,
    PartnerBooking,
    PartnerServiceDuration,
    PartnerWeeklySchedule,
    PartnerScheduleException,
    PartnerScheduleBlock,
    PartnerStaff,
    PartnerWorkOrder,
    PartnerPart,
    PartnerFinancialEntry,
    PartnerPurchaseLine,
    ServiceEstimate,
    PartnerQuoteShare,
    PartnerWorkResource,
    PartnerServicePreset,
    PartnerDailyClose,
    PartnerCustomer,
    PartnerVehicle,
    PartnerSupplier,
    PartnerStockMovement,
    PartnerPayment,
    PartnerTaxInvoice,
    PartnerWorkOrderLine,
)
from .member_display import member_nickname
from .quote_sharing import create_quote_share, estimate_quote_context, work_order_quote_context


def active_staff_for(request):
    if not request.user.is_authenticated:
        return PartnerStaff.objects.none()
    return PartnerStaff.objects.filter(
        user=request.user,
        is_active=True,
        partner__is_active=True,
    ).select_related("partner")


def is_system_admin(request):
    user = request.user
    return bool(user.is_authenticated and user.is_active and user.is_staff and user.is_superuser)


def staff_partners(request, *, booking=False, records=False):
    if is_system_admin(request):
        return Partner.objects.filter(is_active=True).values_list("id", flat=True)
    qs = active_staff_for(request)
    if booking:
        qs = qs.filter(can_manage_bookings=True)
    if records:
        qs = qs.filter(can_verify_records=True)
    return qs.values_list("partner_id", flat=True)


class PartnerAdminSite(AdminSite):
    site_header = "아이러브미니 협력업체 센터"
    site_title = "협력업체 센터"
    index_title = "업체 업무"
    index_template = "core/partner_admin_index.html"
    login_template = "core/partner_admin_login.html"
    site_url = None

    def each_context(self, request):
        context = super().each_context(request)
        context["is_partner_admin"] = True
        context["partner_application_options"] = Partner.objects.filter(is_active=True).order_by("name", "branch_label")
        context["partner_web_push_public_key"] = (
            getattr(settings, "WEB_PUSH_VAPID_PUBLIC_KEY", "") if self.has_permission(request) else ""
        )
        return context

    def has_permission(self, request):
        return bool(
            request.user.is_authenticated
            and request.user.is_active
            and (
                is_system_admin(request)
                or active_staff_for(request).filter(
                    Q(can_manage_bookings=True) | Q(can_verify_records=True)
                ).exists()
            )
        )

    def get_app_list(self, request, app_label=None):
        app_list = super().get_app_list(request, app_label)
        names = {
            "PartnerBooking": "업체 예약",
            "PartnerServiceDuration": "정비 항목별 작업 시간",
            "PartnerWeeklySchedule": "주간 예약 운영시간",
            "PartnerScheduleException": "날짜별 일정 변경·휴무",
            "PartnerScheduleBlock": "임시 예약 차단",
            "PartnerWorkOrder": "정비 접수·작업지시",
            "PartnerWorkResource": "정비사·작업대",
            "PartnerServicePreset": "표준 정비 항목·공임",
            "PartnerDailyClose": "일일 매출 마감",
            "PartnerCustomer": "고객 대장",
            "PartnerVehicle": "업체 차량 대장",
            "PartnerPosIntegration": "POS 연동 설정",
            "PartnerPayment": "POS 결제 내역",
            "PartnerTaxInvoice": "전자세금계산서 발행 준비",
            "PartnerQuoteShare": "견적 링크 관리",
            "LedgerEntry": "인증 정비기록",
        }
        for app in app_list:
            app["name"] = "업체 업무"
            for model in app["models"]:
                model["name"] = names.get(model["object_name"], model["name"])
        return app_list

    def get_urls(self):
        return [
            path(
                "work-orders/<int:work_order_id>/quote/",
                self.admin_view(self.print_work_order_document),
                {"document_type": "quote"},
                name="partner-workorder-quote",
            ),
            path(
                "work-orders/<int:work_order_id>/<str:document_type>/",
                self.admin_view(self.print_work_order_document),
                name="partner-workorder-document",
            ),
            path(
                "estimates/<int:estimate_id>/print/",
                self.admin_view(self.print_service_estimate),
                name="partner-estimate-print",
            ),
            path(
                "quotes/work-orders/<int:work_order_id>/share/",
                self.admin_view(self.share_work_order_quote),
                name="partner-workorder-quote-share",
            ),
            path(
                "quotes/estimates/<int:estimate_id>/share/",
                self.admin_view(self.share_service_estimate),
                name="partner-estimate-quote-share",
            ),
            path(
                "settings/kakao-channel/",
                self.admin_view(self.kakao_channel_settings),
                name="partner-kakao-channel-settings",
            ),
            path(
                "exports/bookings.xlsx",
                self.admin_view(self.export_bookings),
                name="partner-booking-export",
            ),
            path(
                "exports/maintenance-records.xlsx",
                self.admin_view(self.export_maintenance_records),
                name="partner-maintenance-export",
            ),
            path(
                "exports/workshop-data.xlsx",
                self.admin_view(self.export_workshop_data),
                name="partner-workshop-export",
            ),
            path("reports/", self.admin_view(self.operations_report), name="partner-operations-report"),
            path(
                "ledger/<int:entry_id>/evidence/<int:evidence_id>/",
                self.admin_view(self.ledger_evidence_view),
                name="ledger-evidence",
            ),
        ] + super().get_urls()

    @method_decorator(never_cache)
    def index(self, request, extra_context=None):
        context = self.each_context(request)
        context.update(extra_context or {})
        context["title"] = self.index_title
        context["app_list"] = self.get_app_list(request)

        can_manage_work = is_system_admin(request) or active_staff_for(request).filter(can_manage_bookings=True).exists()
        can_verify_records = is_system_admin(request) or active_staff_for(request).filter(can_verify_records=True).exists()
        partner_ids = staff_partners(request, booking=True)
        bookings = PartnerBooking.objects.filter(partner_id__in=partner_ids)
        orders = PartnerWorkOrder.objects.filter(partner_id__in=partner_ids).select_related("partner", "booking", "vehicle")
        partners_without_weekly_hours = Partner.objects.filter(
            pk__in=partner_ids, is_active=True
        ).exclude(
            pk__in=PartnerWeeklySchedule.objects.filter(
                partner_id__in=partner_ids, is_closed=False
            ).values("partner_id")
        ).count()
        today = timezone.localdate()
        active_statuses = [
            PartnerWorkOrder.Status.RECEIVED,
            PartnerWorkOrder.Status.INSPECTING,
            PartnerWorkOrder.Status.QUOTE,
            PartnerWorkOrder.Status.WAITING_APPROVAL,
            PartnerWorkOrder.Status.WAITING_PARTS,
            PartnerWorkOrder.Status.IN_PROGRESS,
            PartnerWorkOrder.Status.READY,
        ]
        context["partner_dashboard"] = {
            "can_manage_work": can_manage_work,
            "can_verify_records": can_verify_records,
            "today_bookings": bookings.filter(scheduled_at__date=today).count(),
            "booking_requests": bookings.filter(status=PartnerBooking.Status.REQUESTED).count(),
            "partners_without_weekly_hours": partners_without_weekly_hours,
            "active_orders": orders.filter(status__in=active_statuses).count(),
            "awaiting_approval": orders.filter(status=PartnerWorkOrder.Status.WAITING_APPROVAL).count(),
            "low_stock_parts": PartnerPart.objects.filter(partner_id__in=partner_ids, is_active=True, stock_quantity__lte=F("reorder_level")).count(),
            "month_income": PartnerFinancialEntry.objects.filter(partner_id__in=partner_ids, kind=PartnerFinancialEntry.Kind.INCOME, occurred_at__year=today.year, occurred_at__month=today.month).aggregate(total=Sum("amount_krw"))["total"] or 0,
            "month_expense": PartnerFinancialEntry.objects.filter(partner_id__in=partner_ids, kind=PartnerFinancialEntry.Kind.EXPENSE, occurred_at__year=today.year, occurred_at__month=today.month).aggregate(total=Sum("amount_krw"))["total"] or 0,
            "recent_orders": orders.order_by("-updated_at", "-pk")[:8],
            "today": today,
        }
        context["kakao_channel_settings_url"] = reverse("partner_admin:partner-kakao-channel-settings")
        # Surface urgent work in one tenant-scoped queue so staff can act without
        # searching across separate booking and repair lists.
        now = timezone.now()
        queue = []
        queue_specs = (
            (PartnerWorkOrder.Status.WAITING_APPROVAL, 0, "고객 승인 대기", "견적 확인·안내"),
            (PartnerWorkOrder.Status.READY, 1, "출고 대기", "수납·차량 인도 확인"),
            (PartnerWorkOrder.Status.WAITING_PARTS, 3, "부품 대기", "입고 일정 확인"),
            (PartnerWorkOrder.Status.RECEIVED, 4, "접수 미처리", "점검 시작"),
        )
        for status, priority, title, instruction in queue_specs:
            for row in orders.filter(status=status).order_by("updated_at", "pk")[:8]:
                queue.append({
                    "priority": priority, "sort_at": row.updated_at, "title": title,
                    "instruction": instruction,
                    "primary": row.plate_number or row.vehicle_model or row.customer_name or "차량 미입력",
                    "secondary": f"{row.customer_name or '고객 미입력'} · {row.updated_at:%m-%d %H:%M} 변경",
                    "url": reverse("partner_admin:core_partnerworkorder_change", args=[row.pk]),
                    "status": row.status,
                })
        stale_bookings = bookings.filter(
            status=PartnerBooking.Status.REQUESTED,
            requested_at__lt=now - timedelta(hours=24),
        ).select_related("vehicle").order_by("requested_at")[:8]
        for booking in stale_bookings:
            queue.append({
                "priority": 2, "sort_at": booking.requested_at,
                "title": "예약 요청 24시간 초과",
                "instruction": "고객에게 예약 가능 여부 회신",
                "primary": booking.vehicle.model_name if booking.vehicle else booking.service_type,
                "secondary": f"{member_nickname(booking.customer)} · {booking.requested_at:%m-%d %H:%M} 요청",
                "url": reverse("partner_admin:core_partnerbooking_change", args=[booking.pk]),
                "status": "booking",
            })
        queue.sort(key=lambda item: (item["priority"], item["sort_at"]))
        context["partner_dashboard"]["action_queue"] = queue[:12]
        return render(request, self.index_template, context)

    def kakao_channel_settings(self, request):
        from urllib.parse import urlparse
        partner_ids = staff_partners(request, booking=True)
        partners = Partner.objects.filter(pk__in=partner_ids, is_active=True).order_by("name", "branch_label")
        selected_id = request.POST.get("partner_id") if request.method == "POST" else request.GET.get("partner")
        partner = partners.filter(pk=selected_id).first() if selected_id else partners.first()
        if not partner:
            raise PermissionDenied
        if request.method == "POST":
            raw_url = request.POST.get("kakao_channel_url", "").strip()
            parsed = urlparse(raw_url) if raw_url else None
            if raw_url and (parsed.scheme != "https" or parsed.hostname not in {"pf.kakao.com", "open.kakao.com"}):
                messages.error(request, "카카오 채널의 HTTPS 주소를 입력해 주세요. (pf.kakao.com 또는 open.kakao.com)")
            else:
                partner.kakao_channel_url = raw_url
                partner.save(update_fields=["kakao_channel_url"])
                messages.success(request, "업체 카카오톡 채널 주소를 저장했습니다.")
                from django.shortcuts import redirect
                return redirect(f"{reverse('partner_admin:partner-kakao-channel-settings')}?partner={partner.pk}")
        return render(request, "core/partner_kakao_channel_settings.html", {
            **self.each_context(request), "title": "카카오톡 채널 설정", "partners": partners,
            "selected_partner": partner,
        })

    def print_work_order_document(self, request, work_order_id, document_type=None, **kwargs):
        document_type = document_type or kwargs.get("document_type")
        document_types = {
            "quote": "정비 견적서",
            "job-card": "정비 작업지시서",
            "repair-statement": "정비 명세서",
            "receipt": "수납 영수증",
        }
        if document_type not in document_types:
            raise Http404
        work_order = get_object_or_404(
            PartnerWorkOrder.objects.select_related("partner", "booking", "vehicle", "customer", "partner_vehicle").prefetch_related("lines__part"),
            pk=work_order_id,
            partner_id__in=staff_partners(request, booking=True),
        )
        context = work_order_quote_context(work_order, document_type)
        if document_type == "quote":
            context["share_endpoint"] = reverse("partner_admin:partner-workorder-quote-share", args=[work_order.pk])
        return render(request, "core/partner_quote_print.html", context)

    def share_work_order_quote(self, request, work_order_id):
        work_order = get_object_or_404(
            PartnerWorkOrder.objects.select_related("partner"), pk=work_order_id,
            partner_id__in=staff_partners(request, booking=True),
        )
        return create_quote_share(request, work_order=work_order)

    def operations_report(self, request):
        today = timezone.localdate()
        requested_month = request.GET.get("month", "")
        try:
            report_month = datetime.strptime(requested_month, "%Y-%m").date().replace(day=1)
        except ValueError:
            report_month = today.replace(day=1)
        partner_ids = staff_partners(request, booking=True)
        income = PartnerFinancialEntry.objects.filter(
            partner_id__in=partner_ids, kind=PartnerFinancialEntry.Kind.INCOME,
            occurred_at__year=report_month.year, occurred_at__month=report_month.month,
        ).aggregate(total=Sum("amount_krw"))["total"] or 0
        expenses = PartnerFinancialEntry.objects.filter(
            partner_id__in=partner_ids, kind=PartnerFinancialEntry.Kind.EXPENSE,
            occurred_at__year=report_month.year, occurred_at__month=report_month.month,
        ).aggregate(total=Sum("amount_krw"))["total"] or 0
        pos_payments = PartnerPayment.objects.filter(
            partner_id__in=partner_ids, status=PartnerPayment.Status.APPROVED,
            paid_at__year=report_month.year, paid_at__month=report_month.month,
        )
        pos_income = pos_payments.aggregate(total=Sum("amount_krw"))["total"] or 0
        pos_refunds = PartnerPayment.objects.filter(
            partner_id__in=partner_ids,
            cancelled_at__year=report_month.year,
            cancelled_at__month=report_month.month,
        ).aggregate(total=Sum("cancelled_amount_krw"))["total"] or 0
        work_orders = PartnerWorkOrder.objects.filter(
            partner_id__in=partner_ids, created_at__year=report_month.year,
            created_at__month=report_month.month,
        ).exclude(status=PartnerWorkOrder.Status.CANCELLED)
        receivable = work_orders.aggregate(total=Sum(F("final_total_krw") - F("paid_amount_krw")))["total"] or 0
        parts = PartnerPart.objects.filter(partner_id__in=partner_ids, is_active=True)
        stock_value = parts.aggregate(total=Sum(F("stock_quantity") * F("purchase_price_krw")))["total"] or 0
        context = self.each_context(request)
        context.update({
            "title": "월간 운영 리포트",
            "report_month": report_month,
            "income": income,
            "expenses": expenses,
            "net_cash": income - expenses,
            "pos_income": pos_income,
            "pos_refunds": pos_refunds,
            "pos_payment_count": pos_payments.count(),
            "receivable": receivable,
            "stock_value": stock_value,
            "completed_orders": work_orders.filter(status=PartnerWorkOrder.Status.COMPLETED).count(),
            "open_orders": work_orders.exclude(status__in=[PartnerWorkOrder.Status.COMPLETED, PartnerWorkOrder.Status.CANCELLED]).count(),
            "recent_finances": PartnerFinancialEntry.objects.filter(
                partner_id__in=partner_ids, occurred_at__year=report_month.year,
                occurred_at__month=report_month.month,
            ).select_related("partner", "work_order")[:20],
            "recent_pos_payments": pos_payments.select_related("partner", "work_order").order_by("-paid_at", "-pk")[:20],
            "recent_purchases": PartnerPurchaseLine.objects.filter(
                partner_id__in=partner_ids, purchased_at__year=report_month.year,
                purchased_at__month=report_month.month,
            ).select_related("partner", "supplier", "part")[:20],
        })
        return render(request, "core/partner_operations_report.html", context)

    def print_service_estimate(self, request, estimate_id):
        estimate = get_object_or_404(
            ServiceEstimate.objects.select_related(
                "booking__partner", "booking__vehicle", "booking__customer__naver_identity",
                "booking__customer__apple_identity",
            ),
            pk=estimate_id,
            booking__partner_id__in=staff_partners(request, booking=True),
        )
        context = estimate_quote_context(estimate)
        context["share_endpoint"] = reverse("partner_admin:partner-estimate-quote-share", args=[estimate.pk])
        return render(request, "core/partner_quote_print.html", context)

    def share_service_estimate(self, request, estimate_id):
        estimate = get_object_or_404(
            ServiceEstimate.objects.select_related("booking__partner"), pk=estimate_id,
            booking__partner_id__in=staff_partners(request, booking=True),
        )
        return create_quote_share(request, estimate=estimate)

    @staticmethod
    def _excel_response(filename, title, headers, rows):
        from io import BytesIO
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = title[:31]
        sheet.append(headers)
        for row in rows:
            safe_row = []
            for value in row:
                if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t"):
                    value = "'" + value
                safe_row.append(value)
            sheet.append(safe_row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(name="맑은 고딕", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="B3213B")
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 26
        for column in sheet.columns:
            values = [len(str(cell.value or "")) for cell in column]
            width = min(max(max(values, default=8) + 2, 12), 42)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = width
            for cell in column[1:]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        response["Cache-Control"] = "private, no-store"
        return response

    def export_bookings(self, request):
        staff = active_staff_for(request).filter(can_manage_bookings=True)
        if not is_system_admin(request) and not staff.exists():
            raise Http404
        partner_ids = staff_partners(request, booking=True)
        bookings = PartnerBooking.objects.filter(partner_id__in=partner_ids).select_related(
            "partner", "customer__naver_identity", "customer__apple_identity", "vehicle"
        ).order_by("-scheduled_at", "-pk")
        rows = []
        for booking in bookings.iterator():
            vehicle = booking.vehicle
            rows.append((
                timezone.localtime(booking.scheduled_at).strftime("%Y-%m-%d %H:%M"),
                vehicle.plate_number if vehicle else "",
                booking.contact_phone or "미등록",
                vehicle.model_name if vehicle else "차량 정보 없음",
                booking.partner.name,
                timezone.localtime(booking.requested_at).strftime("%Y-%m-%d %H:%M"),
                booking.service_type,
                booking.get_status_display(),
                booking.customer_note,
                booking.partner_note,
            ))
        return self._excel_response(
            f"ilovemini-bookings-{timezone.localdate():%Y%m%d}.xlsx",
            "예약 관리대장",
            ["예약 일시", "차량번호", "연락처", "차량", "업체", "요청 일시", "작업 종류", "예약 상태", "요청 사항", "업체 메모"],
            rows,
        )

    def export_maintenance_records(self, request):
        staff = active_staff_for(request).filter(can_verify_records=True)
        if not is_system_admin(request) and not staff.exists():
            raise Http404
        partner_ids = staff_partners(request, records=True)
        entries = LedgerEntry.objects.filter(
            source=LedgerEntry.Source.PARTNER,
            partner_id__in=partner_ids,
        ).select_related("partner", "vehicle", "verified_by__naver_identity", "verified_by__apple_identity").prefetch_related(
            "evidence_images"
        ).order_by("-entry_date", "-pk")
        rows = []
        for entry in entries.iterator(chunk_size=200):
            vehicle = entry.vehicle
            verified_by = entry.verified_by
            line_items = entry.line_items if isinstance(entry.line_items, list) else []
            item_values = []
            for item in line_items:
                if not isinstance(item, dict) or not str(item.get("name", "")).strip():
                    continue
                try:
                    amount = int(item.get("amount_krw") or 0)
                except (TypeError, ValueError):
                    amount = 0
                item_values.append(f"{str(item.get('name')).strip()}: {amount:,}원")
            rows.append((
                entry.entry_date.strftime("%Y-%m-%d"),
                timezone.localtime(entry.verified_at).strftime("%Y-%m-%d %H:%M") if entry.verified_at else "",
                entry.partner.name if entry.partner else "",
                vehicle.model_name if vehicle else "차량 정보 없음",
                vehicle.plate_number if vehicle else "",
                entry.odometer_km,
                entry.get_kind_display(),
                " / ".join(item_values),
                entry.gross_amount_krw if entry.gross_amount_krw is not None else entry.amount_krw,
                entry.amount_krw,
                entry.get_payment_method_display() if entry.payment_method else "미기록(이전 앱)",
                entry.description,
                entry.evidence_images.count(),
                member_nickname(verified_by) if verified_by else "",
            ))
        return self._excel_response(
            f"ilovemini-maintenance-records-{timezone.localdate():%Y%m%d}.xlsx",
            "정비 관리대장",
            ["정비 일자", "등록 일시", "업체", "차량", "차량번호", "주행거리(km)", "작업 종류", "항목별 금액", "정비 총액(포인트 차감 전, 원)", "실제 결제액(원)", "결제 수단", "정비 내용", "첨부 사진 수", "등록 담당자"],
            rows,
        )

    def export_workshop_data(self, request):
        """Download a tenant-scoped workbook of operational data for backup/migration."""
        if not is_system_admin(request) and not active_staff_for(request).filter(can_manage_bookings=True).exists():
            raise Http404
        partner_ids = list(staff_partners(request, booking=True))
        from io import BytesIO
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        workbook.remove(workbook.active)

        def add_sheet(title, headers, queryset, row_builder):
            sheet = workbook.create_sheet(title[:31])
            sheet.append(headers)
            for obj in queryset.iterator(chunk_size=500):
                row = row_builder(obj)
                sheet.append([("'" + value) if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t") else value for value in row])
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.row_dimensions[1].height = 24
            for cell in sheet[1]:
                cell.font = Font(name="맑은 고딕", bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="B3213B")
            for column in sheet.columns:
                width = min(max(max((len(str(cell.value or "")) for cell in column), default=8) + 2, 12), 42)
                sheet.column_dimensions[get_column_letter(column[0].column)].width = width
                for cell in column[1:]:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

        add_sheet("고객", ["업체", "고객명", "연락처", "이메일", "주소", "사업자번호", "메모", "관리 상태", "등록일"],
            PartnerCustomer.objects.filter(partner_id__in=partner_ids).select_related("partner").order_by("partner_id", "pk"),
            lambda x: [x.partner.name, x.name, x.phone, x.email, x.address, x.business_number, x.notes, x.is_active, x.created_at.strftime("%Y-%m-%d")])
        add_sheet("차량", ["업체", "고객", "연락처", "차량번호", "제조사", "차종", "연식", "차대번호", "최초등록일", "주행거리 km", "메모"],
            PartnerVehicle.objects.filter(partner_id__in=partner_ids).select_related("partner", "customer").order_by("partner_id", "pk"),
            lambda x: [x.partner.name, x.customer.name, x.customer.phone, x.plate_number, x.manufacturer, x.model_name, x.model_year or "", x.vin, x.first_registration_date.isoformat() if x.first_registration_date else "", x.current_odometer_km or "", x.notes])
        add_sheet("부품재고", ["업체", "관리번호", "부품번호", "부품명", "제조사", "단위", "위치", "매입단가", "판매단가", "현재고", "재주문 기준", "사용"],
            PartnerPart.objects.filter(partner_id__in=partner_ids).select_related("partner").order_by("partner_id", "pk"),
            lambda x: [x.partner.name, x.sku, x.part_number, x.name, x.maker, x.unit, x.location, x.purchase_price_krw, x.sale_price_krw, x.stock_quantity, x.reorder_level, x.is_active])
        add_sheet("작업지시", ["작업번호", "업체", "접수일", "차량번호", "차량", "고객", "연락처", "주행거리 km", "상태", "견적상태", "견적금액", "최종금액", "수납", "작업자", "요청", "진단", "작업내용"],
            PartnerWorkOrder.objects.filter(partner_id__in=partner_ids).select_related("partner").order_by("partner_id", "pk"),
            lambda x: [f"WO-{x.created_at:%y%m%d}-{x.pk:05d}", x.partner.name, x.created_at.strftime("%Y-%m-%d %H:%M"), x.plate_number, x.vehicle_model, x.customer_name, x.contact_phone, x.odometer_km or "", x.get_status_display(), x.get_quote_status_display(), x.quoted_total_krw, x.final_total_krw, x.paid_amount_krw, x.technician_name, x.customer_concern, x.inspection_notes, x.work_notes])
        add_sheet("정비항목", ["작업번호", "유형", "항목", "수량", "단위", "단가", "금액", "재고 부품"],
            PartnerWorkOrderLine.objects.filter(work_order__partner_id__in=partner_ids).select_related("work_order", "part").order_by("work_order_id", "pk"),
            lambda x: [x.work_order_id, x.get_kind_display(), x.description, str(x.quantity), x.unit, x.unit_price_krw, x.amount_krw, x.part.name if x.part_id else ""])
        add_sheet("매입", ["업체", "매입일", "전표", "공급처", "부품", "수량", "단가", "합계", "지불", "미지불", "메모"],
            PartnerPurchaseLine.objects.filter(partner_id__in=partner_ids).select_related("partner", "supplier", "part").order_by("partner_id", "pk"),
            lambda x: [x.partner.name, x.purchased_at.isoformat(), x.invoice_number, x.supplier.name, x.part.name, x.quantity, x.unit_price_krw, x.total_krw, x.paid_krw, max(0, x.total_krw-x.paid_krw), x.note])
        add_sheet("재고이동", ["처리일", "업체", "부품", "방향", "사유", "수량", "처리후 재고", "작업번호", "메모"],
            PartnerStockMovement.objects.filter(partner_id__in=partner_ids).select_related("partner", "part", "work_order").order_by("partner_id", "pk"),
            lambda x: [x.created_at.strftime("%Y-%m-%d %H:%M"), x.partner.name, x.part.name, x.get_direction_display(), x.get_reason_display(), x.quantity, x.balance_after, x.work_order_id or "", x.note])
        add_sheet("입금경비", ["업체", "처리일", "구분", "항목", "금액", "수단", "상대방", "작업번호", "메모"],
            PartnerFinancialEntry.objects.filter(partner_id__in=partner_ids).select_related("partner", "work_order").order_by("partner_id", "pk"),
            lambda x: [x.partner.name, x.occurred_at.isoformat(), x.get_kind_display(), x.category, x.amount_krw, x.get_payment_method_display() if x.payment_method else "", x.counterparty, x.work_order_id or "", x.note])
        add_sheet("POS결제", ["결제 ID", "업체 ID", "업체", "일시", "취소 일시", "취소 금액", "제공사", "작업 ID", "거래번호", "수단", "상태", "금액", "승인번호"],
            PartnerPayment.objects.filter(partner_id__in=partner_ids).select_related("partner").order_by("partner_id", "pk"),
            lambda x: [x.pk, x.partner_id, x.partner.name, x.paid_at.strftime("%Y-%m-%d %H:%M"), x.cancelled_at.strftime("%Y-%m-%d %H:%M") if x.cancelled_at else "", x.cancelled_amount_krw, x.provider_name, x.work_order_id, x.provider_transaction_id, x.get_method_display(), x.get_status_display(), x.amount_krw, x.approval_code])
        add_sheet("세금계산서", ["업체", "작성일", "상태", "유형", "목적", "공급자 사업자번호", "공급자", "공급받는자 사업자번호", "공급받는자", "공급가", "세액", "합계", "외부문서번호"],
            PartnerTaxInvoice.objects.filter(partner_id__in=partner_ids).select_related("partner").order_by("partner_id", "pk"),
            lambda x: [x.partner.name, x.write_date.isoformat(), x.get_status_display(), x.get_tax_type_display(), x.get_purpose_display(), x.seller_business_number, x.seller_name, x.buyer_business_number, x.buyer_name, x.supply_cost_krw, x.tax_krw, x.total_krw, x.provider_document_id])
        add_sheet("공급처", ["업체", "공급처", "담당자", "연락처", "이메일", "사업자번호", "주소", "메모", "사용"],
            PartnerSupplier.objects.filter(partner_id__in=partner_ids).select_related("partner").order_by("partner_id", "pk"),
            lambda x: [x.partner.name, x.name, x.contact_name, x.phone, x.email, x.business_number, x.address, x.note, x.is_active])
        add_sheet("정비템플릿", ["업체", "항목", "유형", "소요시간(분)", "기본수량", "단위", "기본단가", "연결부품", "사용"],
            PartnerServicePreset.objects.filter(partner_id__in=partner_ids).select_related("partner", "part").order_by("partner_id", "pk"),
            lambda x: [x.partner.name, x.name, x.get_kind_display(), x.duration_minutes, str(x.default_quantity), x.unit, x.default_price_krw, x.part.name if x.part_id else "", x.is_active])
        add_sheet("작업자작업대", ["업체", "구분", "이름", "사용"],
            PartnerWorkResource.objects.filter(partner_id__in=partner_ids).select_related("partner").order_by("partner_id", "pk"),
            lambda x: [x.partner.name, x.get_kind_display(), x.name, x.is_active])
        add_sheet("일일마감", ["마감 ID", "업체 ID", "업체", "영업일", "POS 현금", "실제 현금", "현금 차이", "POS 카드", "실제 카드", "카드 차이", "POS 이체", "실제 이체", "이체 차이", "POS 환불", "실제 환불", "환불 차이", "메모"],
            PartnerDailyClose.objects.filter(partner_id__in=partner_ids).select_related("partner").order_by("partner_id", "business_date"),
            lambda x: [x.pk, x.partner_id, x.partner.name, x.business_date.isoformat(), x.expected_cash_krw, x.counted_cash_krw, x.counted_cash_krw-x.expected_cash_krw, x.expected_card_krw, x.counted_card_krw, x.counted_card_krw-x.expected_card_krw, x.expected_transfer_krw, x.counted_transfer_krw, x.counted_transfer_krw-x.expected_transfer_krw, x.expected_refund_krw, x.counted_refund_krw, x.counted_refund_krw-x.expected_refund_krw, x.note])

        # Include the booking rules and reservations required to reconstruct a
        # workshop calendar. Row IDs and foreign keys are stable restore keys.
        add_sheet("예약", ["예약 ID", "업체 ID", "고객 계정 ID", "차량 ID", "시작 시각", "소요 분", "담당 자원 ID", "정비 종류", "상태", "연락처", "고객 요청", "업체 메모"],
            PartnerBooking.objects.filter(partner_id__in=partner_ids).order_by("partner_id", "pk"),
            lambda x: [x.pk, x.partner_id, x.customer_id, x.vehicle_id or "", timezone.localtime(x.scheduled_at).isoformat(), x.duration_minutes, x.assigned_resource_id or "", x.service_type, x.get_status_display(), x.contact_phone, x.customer_note, x.partner_note])
        add_sheet("주간일정", ["일정 ID", "업체 ID", "요일", "개점", "폐점", "예약 간격(분)", "휴무"],
            PartnerWeeklySchedule.objects.filter(partner_id__in=partner_ids).order_by("partner_id", "weekday"),
            lambda x: [x.pk, x.partner_id, x.weekday, x.opens_at.isoformat(), x.closes_at.isoformat(), x.slot_interval_minutes, x.is_closed])
        add_sheet("일정 예외", ["예외 ID", "업체 ID", "날짜", "휴무", "개점", "폐점", "메모"],
            PartnerScheduleException.objects.filter(partner_id__in=partner_ids).order_by("partner_id", "date"),
            lambda x: [x.pk, x.partner_id, x.date.isoformat(), x.is_closed, x.opens_at.isoformat() if x.opens_at else "", x.closes_at.isoformat() if x.closes_at else "", x.note])
        add_sheet("예약 차단", ["차단 ID", "업체 ID", "시작", "종료", "사유"],
            PartnerScheduleBlock.objects.filter(partner_id__in=partner_ids).order_by("partner_id", "starts_at"),
            lambda x: [x.pk, x.partner_id, timezone.localtime(x.starts_at).isoformat(), timezone.localtime(x.ends_at).isoformat(), x.reason])
        add_sheet("정비 소요시간", ["설정 ID", "업체 ID", "정비 종류", "소요 분", "사용"],
            PartnerServiceDuration.objects.filter(partner_id__in=partner_ids).order_by("partner_id", "service_name"),
            lambda x: [x.pk, x.partner_id, x.service_name, x.duration_minutes, x.is_active])

        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="ilovemini-workshop-backup-{timezone.localdate():%Y%m%d}.xlsx"'
        response["Cache-Control"] = "private, no-store"
        return response

    def ledger_evidence_view(self, request, entry_id, evidence_id):
        if not is_system_admin(request) and not active_staff_for(request).filter(can_verify_records=True).exists():
            raise Http404
        partner_ids = staff_partners(request, records=True)
        evidence = LedgerEvidence.objects.select_related("entry").filter(
            pk=evidence_id,
            entry_id=entry_id,
            entry__source=LedgerEntry.Source.PARTNER,
            entry__partner_id__in=partner_ids,
        ).first()
        if not evidence:
            raise Http404
        response = HttpResponse(bytes(evidence.image_data), content_type=evidence.content_type)
        response["Cache-Control"] = "private, no-store"
        response["X-Content-Type-Options"] = "nosniff"
        return response


partner_admin_site = PartnerAdminSite(name="partner_admin")


@admin.register(PartnerQuoteShare, site=partner_admin_site)
class PartnerQuoteShareAdmin(admin.ModelAdmin):
    list_display = ("token", "source_label", "quote_revision", "created_at", "expires_at", "revoked_at")
    list_filter = ("created_at", "expires_at", "revoked_at")
    readonly_fields = ("token", "work_order", "estimate", "quote_revision", "created_by", "created_at", "expires_at")
    fields = readonly_fields + ("revoked_at",)
    actions = ("revoke_selected",)

    def has_add_permission(self, request):
        return False

    def has_delete_permission(self, request, obj=None):
        return False

    def get_queryset(self, request):
        qs = super().get_queryset(request).select_related("work_order__partner", "estimate__booking__partner")
        partner_ids = staff_partners(request, booking=True)
        return qs.filter(Q(work_order__partner_id__in=partner_ids) | Q(estimate__booking__partner_id__in=partner_ids))

    @admin.display(description="견적")
    def source_label(self, obj):
        return f"{obj.work_order or obj.estimate}"

    @admin.action(description="선택한 공유 링크 만료 처리")
    def revoke_selected(self, request, queryset):
        queryset.filter(revoked_at__isnull=True).update(revoked_at=timezone.now())


class PartnerBookingStatusFilter(admin.SimpleListFilter):
    title = "예약 상태"
    parameter_name = "status"

    def lookups(self, request, model_admin):
        return PartnerBooking.Status.choices

    def queryset(self, request, queryset):
        return queryset.filter(status=self.value()) if self.value() else queryset


class PartnerBookingAdminForm(forms.ModelForm):
    class Meta:
        model = PartnerBooking
        fields = "__all__"

    def clean(self):
        data = super().clean()
        resource = data.get("assigned_resource")
        if resource and self.instance.partner_id and resource.partner_id != self.instance.partner_id:
            self.add_error("assigned_resource", "해당 예약 업체에 등록된 정비사·작업대만 배정할 수 있습니다.")
        return data


@admin.register(PartnerBooking, site=partner_admin_site)
class PartnerBookingAdmin(admin.ModelAdmin):
    form = PartnerBookingAdminForm
    change_list_template = "admin/core/partnerbooking/change_list.html"
    change_form_template = "admin/core/partnerbooking/change_form.html"
    list_display = ("booking_time", "vehicle_plate", "contact_phone_display", "vehicle_model", "service_label", "duration_label", "schedule_conflict_label", "status_label")
    list_filter = (PartnerBookingStatusFilter,)
    search_fields = ("vehicle__model_name", "vehicle__plate_number", "contact_phone", "service_type")
    list_select_related = ("vehicle", "partner")
    date_hierarchy = None
    ordering = ("-requested_at",)
    actions = None
    fields = ("vehicle_registration", "contact_phone_display", "booking_time", "scheduled_at", "service_label", "duration_minutes", "assigned_resource", "customer_note_display", "status", "partner_note")
    readonly_fields = ("vehicle_registration", "contact_phone_display", "booking_time", "service_label", "customer_note_display")

    def formfield_for_foreignkey(self, db_field, request, **kwargs):
        if db_field.name == "assigned_resource":
            kwargs["queryset"] = PartnerWorkResource.objects.filter(
                partner_id__in=staff_partners(request, booking=True), is_active=True,
            )
        return super().formfield_for_foreignkey(db_field, request, **kwargs)

    @admin.display(description="차량번호 / 차량")
    def vehicle_registration(self, obj):
        if not obj.vehicle:
            return "차량 정보 없음"
        plate = obj.vehicle.plate_number or "차량번호 미등록"
        return f"{plate} · {obj.vehicle.model_name}"

    @admin.display(description="예약 일시", ordering="scheduled_at")
    def booking_time(self, obj):
        return timezone.localtime(obj.scheduled_at).strftime("%Y-%m-%d %H:%M")

    @admin.display(description="차량번호", ordering="vehicle__plate_number")
    def vehicle_plate(self, obj):
        return obj.vehicle.plate_number if obj.vehicle and obj.vehicle.plate_number else "미등록"

    @admin.display(description="차량")
    def vehicle_model(self, obj):
        return obj.vehicle.model_name if obj.vehicle else "차량 정보 없음"

    @admin.display(description="연락처", ordering="contact_phone")
    def contact_phone_display(self, obj):
        return obj.contact_phone or "미등록"

    @admin.display(description="작업 종류", ordering="service_type")
    def service_label(self, obj):
        return obj.service_type or "미등록"

    @admin.display(description="예상 작업 시간", ordering="duration_minutes")
    def duration_label(self, obj):
        hours, minutes = divmod(obj.duration_minutes or 60, 60)
        return f"{hours}시간" + (f" {minutes}분" if minutes else "")

    @admin.display(description="일정 확인")
    def schedule_conflict_label(self, obj):
        from .booking_schedule import ACTIVE_BOOKING_STATUSES
        start = obj.scheduled_at
        end = start + timedelta(minutes=obj.duration_minutes or 60)
        conflicts = PartnerBooking.objects.filter(partner=obj.partner, status__in=ACTIVE_BOOKING_STATUSES, scheduled_at__lt=end).exclude(pk=obj.pk)
        if obj.assigned_resource_id:
            conflicts = conflicts.filter(Q(assigned_resource__isnull=True) | Q(assigned_resource_id=obj.assigned_resource_id))
        for other in conflicts.only("scheduled_at", "duration_minutes"):
            if other.scheduled_at + timedelta(minutes=other.duration_minutes or 60) > start:
                return "재조정 필요"
        return "정상"

    @admin.display(description="예약 상태", ordering="status")
    def status_label(self, obj):
        return obj.get_status_display()

    @admin.display(description="요청 사항")
    def customer_note_display(self, obj):
        return obj.customer_note or "없음"

    def formfield_for_dbfield(self, db_field, request, **kwargs):
        field = super().formfield_for_dbfield(db_field, request, **kwargs)
        labels = {"status": "예약 상태", "partner_note": "업체 메모", "duration_minutes": "예상 작업 시간(분)"}
        if field and db_field.name in labels:
            field.label = labels[db_field.name]
        if field and db_field.name == "status":
            field.widget.attrs.update({"class": "partner-booking-status"})
        if field and db_field.name == "duration_minutes":
            choices = [(value, str(value // 60) + "시간" + (" " + str(value % 60) + "분" if value % 60 else "")) for value in range(30, 481, 30)]
            field.widget = forms.Select(choices=choices)
        return field

    def get_queryset(self, request):
        return super().get_queryset(request).filter(
            partner_id__in=staff_partners(request, booking=True)
        ).select_related("partner", "customer", "vehicle", "assigned_resource")

    def changelist_view(self, request, extra_context=None):
        extra_context = extra_context or {}
        extra_context["title"] = "업체 예약"
        today = timezone.localdate()
        year = request.GET.get("scheduled_at__year", "")
        month_number = request.GET.get("scheduled_at__month", "")
        try:
            month_start = datetime.strptime(f"{year}-{month_number}", "%Y-%m").date().replace(day=1)
        except ValueError:
            month_start = today.replace(day=1)
        next_month = (month_start.replace(day=28) + timedelta(days=4)).replace(day=1)
        prev_month = (month_start - timedelta(days=1)).replace(day=1)
        grid_start = month_start - timedelta(days=month_start.weekday())
        final_day = next_month - timedelta(days=1)
        grid_end = final_day + timedelta(days=6 - final_day.weekday())
        bookings = self.get_queryset(request).filter(
            scheduled_at__gte=timezone.make_aware(datetime.combine(month_start, time.min)),
            scheduled_at__lt=timezone.make_aware(datetime.combine(next_month, time.min)),
        ).order_by("scheduled_at", "pk")
        events_by_day = {}
        booking_rows = list(bookings)
        for booking in booking_rows:
            vehicle = booking.vehicle
            start = booking.scheduled_at
            end = start + timedelta(minutes=booking.duration_minutes or 60)
            has_conflict = any(
                other.pk != booking.pk and other.partner_id == booking.partner_id
                and (not booking.assigned_resource_id or not other.assigned_resource_id or booking.assigned_resource_id == other.assigned_resource_id)
                and other.scheduled_at < end
                and other.scheduled_at + timedelta(minutes=other.duration_minutes or 60) > start
                and other.status in (PartnerBooking.Status.REQUESTED, PartnerBooking.Status.CONFIRMED)
                and booking.status in (PartnerBooking.Status.REQUESTED, PartnerBooking.Status.CONFIRMED)
                for other in booking_rows
            )
            event = {
                "time": timezone.localtime(booking.scheduled_at).strftime("%H:%M"),
                "ends_at": timezone.localtime(end).strftime("%H:%M"),
                "has_conflict": has_conflict,
                "plate": vehicle.plate_number if vehicle and vehicle.plate_number else "차량번호 미등록",
                "phone": booking.contact_phone or "연락처 미등록",
                "service": booking.service_type,
                "resource": str(booking.assigned_resource) if booking.assigned_resource_id else "미배정",
                "status": booking.get_status_display(),
                "status_code": booking.status,
                "change_url": reverse(f"{self.admin_site.name}:core_partnerbooking_change", args=[booking.pk]) if booking.status in (PartnerBooking.Status.REQUESTED, PartnerBooking.Status.CONFIRMED) else "",
            }
            events_by_day.setdefault(timezone.localtime(booking.scheduled_at).date(), []).append(event)
        weeks = []
        current_day = grid_start
        while current_day <= grid_end:
            week = []
            for _ in range(7):
                week.append({
                    "date": current_day,
                    "day": current_day.day,
                    "in_month": current_day.month == month_start.month,
                    "is_today": current_day == today,
                    "events": events_by_day.get(current_day, []),
                })
                current_day += timedelta(days=1)
            weeks.append(week)
        extra_context.update({
            "partner_calendar_month": month_start,
            "partner_calendar_prev": prev_month.strftime("%Y-%m"),
            "partner_calendar_next": next_month.strftime("%Y-%m"),
            "partner_calendar_weeks": weeks,
            "partner_calendar_weekdays": ["월", "화", "수", "목", "금", "토", "일"],
        })
        return super().changelist_view(request, extra_context=extra_context)

    def has_module_permission(self, request):
        return is_system_admin(request) or active_staff_for(request).filter(can_manage_bookings=True).exists()

    def has_view_permission(self, request, obj=None):
        allowed = self.has_module_permission(request)
        return allowed and (obj is None or obj.partner_id in staff_partners(request, booking=True))

    def has_add_permission(self, request):
        return False

    def has_change_permission(self, request, obj=None):
        allowed = self.has_module_permission(request)
        if not allowed or (obj is not None and obj.partner_id not in staff_partners(request, booking=True)):
            return False
        return obj is None or obj.status in (PartnerBooking.Status.REQUESTED, PartnerBooking.Status.CONFIRMED)

    def has_delete_permission(self, request, obj=None):
        return False

    def save_model(self, request, obj, form, change):
        if not change or obj.partner_id not in staff_partners(request, booking=True):
            raise Http404
        previous = PartnerBooking.objects.filter(pk=obj.pk, partner_id=obj.partner_id).values("status", "partner_note", "scheduled_at", "duration_minutes", "assigned_resource_id").first()
        if previous is None:
            raise Http404
        allowed = {
            PartnerBooking.Status.REQUESTED: {PartnerBooking.Status.CONFIRMED, PartnerBooking.Status.REJECTED},
            PartnerBooking.Status.CONFIRMED: {PartnerBooking.Status.COMPLETED},
        }
        if obj.status != previous["status"] and obj.status not in allowed.get(previous["status"], set()):
            raise PermissionDenied("현재 상태에서 선택할 수 없는 예약 변경입니다.")
        # Capacity matters while a booking still occupies a live slot. A
        # conflicting request must remain rejectable, and completed bookings
        # must not be blocked by another reservation's current schedule.
        from .booking_schedule import ACTIVE_BOOKING_STATUSES, is_slot_available
        if obj.status in ACTIVE_BOOKING_STATUSES:
            available, reason = is_slot_available(
                obj.partner, obj.scheduled_at, obj.duration_minutes,
                exclude_booking_id=obj.pk, resource_id=obj.assigned_resource_id,
            )
            duration_extension = (
                obj.duration_minutes > previous["duration_minutes"]
                and obj.scheduled_at == previous["scheduled_at"]
            )
            if not available and not duration_extension:
                raise ValidationError({"duration_minutes": reason})
        super().save_model(request, obj, form, change)
        if obj.status != previous["status"]:
            def notify_customer():
                try:
                    from .push import send_booking_status_notification
                    send_booking_status_notification(obj)
                except Exception:
                    import logging
                    logging.getLogger(__name__).exception("Partner admin booking notification failed (booking_id=%s)", obj.pk)
            transaction.on_commit(notify_customer)
        elif obj.scheduled_at != previous["scheduled_at"] or obj.duration_minutes != previous["duration_minutes"]:
            def notify_customer_schedule():
                try:
                    from .push import send_booking_schedule_notification
                    send_booking_schedule_notification(obj)
                except Exception:
                    import logging
                    logging.getLogger(__name__).exception("Partner admin booking schedule notification failed (booking_id=%s)", obj.pk)
            transaction.on_commit(notify_customer_schedule)


class PartnerScopedScheduleAdmin(admin.ModelAdmin):
    """Keep workshop schedule controls inside the signed-in staff member's scope."""
    booking_scope = True
    list_filter = ()
    search_fields = ("partner__name", "partner__branch_label")

    def get_queryset(self, request):
        return super().get_queryset(request).filter(partner_id__in=staff_partners(request, booking=self.booking_scope))

    def has_module_permission(self, request):
        return active_staff_for(request).filter(can_manage_bookings=True).exists() or is_system_admin(request)

    def has_view_permission(self, request, obj=None):
        return self.has_module_permission(request) and (obj is None or obj.partner_id in staff_partners(request, booking=self.booking_scope))

    def has_change_permission(self, request, obj=None):
        return self.has_view_permission(request, obj)

    def has_add_permission(self, request):
        return self.has_module_permission(request)

    def has_delete_permission(self, request, obj=None):
        return self.has_view_permission(request, obj)

    def formfield_for_foreignkey(self, db_field, request, **kwargs):
        field = super().formfield_for_foreignkey(db_field, request, **kwargs)
        if db_field.name == "partner" and field:
            field.queryset = Partner.objects.filter(pk__in=staff_partners(request, booking=self.booking_scope))
        return field


@admin.register(PartnerServiceDuration, site=partner_admin_site)
class PartnerServiceDurationAdmin(PartnerScopedScheduleAdmin):
    list_display = ("partner", "service_name", "duration_minutes", "is_active")
    list_editable = ("duration_minutes", "is_active")
    fields = ("partner", "service_name", "duration_minutes", "is_active")


@admin.register(PartnerWeeklySchedule, site=partner_admin_site)
class PartnerWeeklyScheduleAdmin(PartnerScopedScheduleAdmin):
    list_display = ("partner", "weekday", "opens_at", "closes_at", "slot_interval_minutes", "is_closed")
    list_editable = ("opens_at", "closes_at", "slot_interval_minutes", "is_closed")
    fields = ("partner", "weekday", "opens_at", "closes_at", "slot_interval_minutes", "is_closed")


@admin.register(PartnerScheduleException, site=partner_admin_site)
class PartnerScheduleExceptionAdmin(PartnerScopedScheduleAdmin):
    list_display = ("partner", "date", "is_closed", "opens_at", "closes_at", "note")
    list_editable = ("is_closed", "opens_at", "closes_at", "note")
    fields = ("partner", "date", "is_closed", "opens_at", "closes_at", "note")


@admin.register(PartnerScheduleBlock, site=partner_admin_site)
class PartnerScheduleBlockAdmin(PartnerScopedScheduleAdmin):
    list_display = ("partner", "starts_at", "ends_at", "reason")
    list_editable = ("ends_at", "reason")
    fields = ("partner", "starts_at", "ends_at", "reason")


@admin.register(LedgerEntry, site=partner_admin_site)
class PartnerLedgerEntryAdmin(admin.ModelAdmin):
    change_list_template = "admin/core/ledgerentry/change_list.html"
    list_display = ("entry_date", "vehicle_plate", "vehicle_model", "kind_display", "amount_display", "odometer_display", "description_display", "evidence_count")
    list_filter = ("kind", "entry_date")
    search_fields = ("vehicle__model_name", "vehicle__plate_number", "description")
    list_select_related = ("partner", "vehicle")
    date_hierarchy = "entry_date"
    ordering = ("-entry_date", "-pk")
    actions = None

    def changelist_view(self, request, extra_context=None):
        extra_context = extra_context or {}
        extra_context.setdefault("title", "인증 정비기록 조회")
        today = timezone.localdate()
        year = request.GET.get("entry_date__year", "")
        month_number = request.GET.get("entry_date__month", "")
        try:
            month_start = datetime.strptime(f"{year}-{month_number}", "%Y-%m").date().replace(day=1)
        except ValueError:
            month_start = today.replace(day=1)
        next_month = (month_start.replace(day=28) + timedelta(days=4)).replace(day=1)
        prev_month = (month_start - timedelta(days=1)).replace(day=1)
        grid_start = month_start - timedelta(days=month_start.weekday())
        final_day = next_month - timedelta(days=1)
        grid_end = final_day + timedelta(days=6 - final_day.weekday())
        entries = self.get_queryset(request).filter(
            entry_date__gte=month_start, entry_date__lt=next_month
        ).order_by("entry_date", "pk")
        entries_by_day = {}
        for entry in entries:
            vehicle = entry.vehicle
            entries_by_day.setdefault(entry.entry_date, []).append({
                "plate": vehicle.plate_number if vehicle and vehicle.plate_number else "번호 미등록",
                "kind": entry.get_kind_display(),
                "amount": f"{entry.amount_krw:,}원",
            })
        weeks = []
        current_day = grid_start
        while current_day <= grid_end:
            week = []
            for _ in range(7):
                week.append({
                    "date": current_day,
                    "day": current_day.day,
                    "in_month": current_day.month == month_start.month,
                    "is_today": current_day == today,
                    "events": entries_by_day.get(current_day, []),
                })
                current_day += timedelta(days=1)
            weeks.append(week)
        extra_context.update({
            "partner_calendar_month": month_start,
            "partner_calendar_prev": prev_month.strftime("%Y-%m"),
            "partner_calendar_next": next_month.strftime("%Y-%m"),
            "partner_calendar_weeks": weeks,
            "partner_calendar_weekdays": ["월", "화", "수", "목", "금", "토", "일"],
        })
        return super().changelist_view(request, extra_context=extra_context)

    @admin.display(description="차량번호", ordering="vehicle__plate_number")
    def vehicle_plate(self, obj):
        return obj.vehicle.plate_number if obj.vehicle and obj.vehicle.plate_number else "미등록"

    @admin.display(description="차량", ordering="vehicle__model_name")
    def vehicle_model(self, obj):
        return obj.vehicle.model_name if obj.vehicle else "차량 정보 없음"

    @admin.display(description="작업 종류", ordering="kind")
    def kind_display(self, obj):
        return obj.get_kind_display()

    @admin.display(description="총액(원)", ordering="amount_krw")
    def amount_display(self, obj):
        return f"{obj.amount_krw:,}"

    @admin.display(description="주행거리(km)", ordering="odometer_km")
    def odometer_display(self, obj):
        return obj.odometer_km if obj.odometer_km is not None else "미등록"

    @admin.display(description="정비 내용", ordering="description")
    def description_display(self, obj):
        return obj.description or "없음"

    def get_queryset(self, request):
        return super().get_queryset(request).filter(
            source=LedgerEntry.Source.PARTNER,
            partner_id__in=staff_partners(request, records=True),
        ).select_related("partner", "vehicle")

    def has_module_permission(self, request):
        return is_system_admin(request) or active_staff_for(request).filter(can_verify_records=True).exists()

    def has_view_permission(self, request, obj=None):
        allowed = self.has_module_permission(request)
        return allowed and (obj is None or obj.partner_id in staff_partners(request, records=True) and obj.source == LedgerEntry.Source.PARTNER)

    def has_add_permission(self, request): return False
    def has_change_permission(self, request, obj=None): return False
    def has_delete_permission(self, request, obj=None): return False

    @admin.display(description="첨부 사진")
    def evidence_count(self, obj):
        images = obj.evidence_images.all()
        if not images:
            return "없음"
        url_name = f"{self.admin_site.name}:ledger-evidence"
        return format_html_join(
            " ", '<a href="{}" target="_blank" rel="noopener">사진 {}</a>',
            (
                (reverse(url_name, kwargs={"entry_id": obj.pk, "evidence_id": image.pk}), index)
                for index, image in enumerate(images, start=1)
            ),
        )

from .models import ServiceEstimate, MaintenanceRecommendation
from .phase_one_admin import EstimateAdmin, RecommendationAdmin
partner_admin_site.register(ServiceEstimate, EstimateAdmin)
partner_admin_site.register(MaintenanceRecommendation, RecommendationAdmin)
from .models import PartnerWorkOrder
from .work_order_admin import PartnerWorkOrderAdmin
partner_admin_site.register(PartnerWorkOrder, PartnerWorkOrderAdmin)
from . import partner_operations_admin  # noqa: F401
